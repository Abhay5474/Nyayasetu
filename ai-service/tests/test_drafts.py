import datetime as dt
import io
import shutil

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pydantic import TypeAdapter, ValidationError

from app.drafts import render, templates
from app.drafts.schemas import DraftRequest, FeePaymentReply, FirstAppeal, RtiApplication
from app.main import app

H = {"X-Internal-Token": "test-token"}
TODAY = dt.date.today()


def rti(**kw):
    base = dict(kind="rti_application", language="en", full_name="Ramesh Kulkarni", address="Kothrud, Pune 411038", place="Pune",
                authority="Pune Municipal Corporation, Shivajinagar, Pune 411005", subject="Road repair",
                items=["Copies of the work orders for Ward 12", "Amount paid to each contractor"])
    return RtiApplication(**{**base, **kw})


def appeal(**kw):
    base = dict(kind="first_appeal", language="en", full_name="Ramesh Kulkarni", address="Kothrud, Pune", place="Pune",
                appellate_authority="First Appellate Authority, Pune Municipal Corporation", pio_details="Shri A. B. Patil, SPIO",
                application_date=TODAY - dt.timedelta(days=70), grounds=["no_reply"], info_subject="Road repair work orders",
                info_office="Road Department")
    return FirstAppeal(**{**base, **kw})


def fee(**kw):
    base = dict(kind="fee_payment_reply", language="en", full_name="Ramesh Kulkarni", address="Kothrud, Pune", place="Pune",
                pio_office="Pune Municipal Corporation", pio_letter_ref="PMC/1187", pio_letter_date=TODAY - dt.timedelta(days=10),
                application_date=TODAY - dt.timedelta(days=30), amount_demanded=148, pages=74, amount_paid=148,
                payment_mode="demand_draft", payment_reference="DD 5523", payment_date=TODAY - dt.timedelta(days=5))
    return FeePaymentReply(**{**base, **kw})


def text_of(doc):
    out = []
    for b in doc.blocks:
        out += [b.text, *b.items, *(f"{k} {v}" for k, v in b.rows)]
    return "\n".join(out)


# ----------------------------------------------------------------------------- RTI application
def test_rti_application_follows_annexure_a():
    doc = templates.build(rti(phone="9876543210", period="1 April 2024 to 31 March 2025"))
    t = text_of(doc)
    assert "Annexure \"A\"" in t and "State Public Information Officer" in t and "Pune Municipal Corporation" in t
    assert "1. Copies of the work orders" in t and "2. Amount paid" in t
    assert "Rs. 10/-" in t and "court fee stamp" in t and "section 7(1)".lower() in t.lower()
    assert doc.blocks[0].type == "box" and "Rs. 10/-" in doc.blocks[0].text  # stamp box
    assert doc.filename.startswith("RTI-application-")


def test_rti_application_marathi_uses_devanagari_numerals():
    doc = templates.build(rti(language="mr", full_name="रमेश कुलकर्णी", address="कोथरूड, पुणे", authority="पुणे महानगरपालिका", date=dt.date(2026, 10, 14)))
    t = text_of(doc)
    assert "१४/१०/२०२६" in t and "१. Copies" in t and "रुपये १०/-" in t and "जोडपत्र “अ”" in t
    assert not any(ch in t.split("Copies")[0] for ch in "0123456789")  # no Western digits in the template's own text


def test_rti_application_bpl_has_no_fee_paragraph():
    doc = templates.build(rti(below_poverty_line=True, bpl_card_no="BPL-123"))
    t = text_of(doc)
    assert "below the poverty line" in t and "BPL-123" in t and "application fee" not in t
    assert not any(b.type == "box" for b in doc.blocks)


def test_rti_application_fee_modes_and_validation():
    t = text_of(templates.build(rti(fee_mode="demand_draft", fee_reference="DD 991122")))
    assert "demand draft no. DD 991122" in t
    with pytest.raises(ValidationError, match="fee_reference"):
        rti(fee_mode="cash")
    with pytest.raises(ValidationError):
        rti(items=[])
    with pytest.raises(ValidationError):
        rti(date=TODAY + dt.timedelta(days=90))


def test_rti_application_warns_when_longer_than_150_words():
    doc = templates.build(rti(items=["word " * 20] * 10))  # 200 words
    assert any("150" in w for w in doc.warnings)
    assert not templates.build(rti()).warnings


def test_rti_application_post_delivery_adds_postage_note():
    t = text_of(templates.build(rti(delivery="post", post_mode="registered")))
    assert "Registered" in t and "postal charges" in t


# ----------------------------------------------------------------------------- first appeal
def test_appeal_deadline_after_a_decision_and_after_no_decision():
    with_order = appeal(order_received_date=TODAY - dt.timedelta(days=5), grounds=["incomplete"])
    assert templates.appeal_last_date(with_order) == TODAY + dt.timedelta(days=25)
    no_reply = appeal(application_date=dt.date(2026, 1, 10))
    assert templates.appeal_last_date(no_reply) == dt.date(2026, 3, 11)  # 30 days for the PIO + 30 days to appeal


def test_appeal_late_adds_condonation_request_and_warns_without_reason():
    doc = templates.build(appeal())
    assert any("ended on" in w for w in doc.warnings) and "Request to condone delay" in text_of(doc) and "[state the reason]" in text_of(doc)
    doc = templates.build(appeal(delay_reason="I was in hospital."))
    assert not doc.warnings and "I was in hospital." in text_of(doc)


def test_appeal_in_time_has_no_condonation_and_cites_provisions():
    doc = templates.build(appeal(application_date=TODAY - dt.timedelta(days=40), grounds=["no_reply", "other"], grounds_other="Extra ground."))
    t = text_of(doc)
    assert "Request to condone delay" not in t and "Section 7(2)" in t and "Section 7(6)" in t and "Extra ground." in t
    assert "Rs. 20/-" in t and "Section 19(1)" in t
    assert any("Section 20" in n for n in doc.notes)  # the note says a penalty is for the Commission, not this appeal
    assert not any("penalty" in b.text.lower() for b in doc.blocks if b.type == "numbered")  # the FAA cannot impose one


def test_appeal_premature_and_inconsistent_grounds_are_flagged():
    early = templates.build(appeal(application_date=TODAY - dt.timedelta(days=10)))
    assert any("still has time" in w for w in early.warnings)
    mismatch = templates.build(appeal(order_received_date=TODAY - dt.timedelta(days=3), application_date=TODAY - dt.timedelta(days=20), grounds=["no_reply"]))
    assert any("Remove one" in w for w in mismatch.warnings)
    missing = templates.build(appeal(application_date=TODAY - dt.timedelta(days=40), grounds=["incomplete"]))
    assert any("no reply" in w for w in missing.warnings)


def test_appeal_validation():
    with pytest.raises(ValidationError, match="grounds_other"):
        appeal(grounds=["other"])
    with pytest.raises(ValidationError, match="before your application"):
        appeal(order_received_date=TODAY - dt.timedelta(days=90))
    with pytest.raises(ValidationError):
        appeal(application_date=TODAY + dt.timedelta(days=2))


# ----------------------------------------------------------------------------- fee payment reply
def test_fee_reply_asks_for_calculation_only_when_the_amount_looks_wrong():
    ok = templates.build(fee())  # 74 pages x Rs. 2 = 148
    assert "calculation" not in text_of(ok) and not ok.warnings
    bad = templates.build(fee(amount_demanded=300, amount_paid=300))
    assert "Rs. 148/-" in text_of(bad) and "calculation of the fee" in text_of(bad) and "7(3)(b)" in text_of(bad) and bad.warnings


def test_fee_reply_warns_on_underpayment():
    assert any("less than" in w for w in templates.build(fee(amount_paid=100)).warnings)
    assert any("more than" in w for w in templates.build(fee(amount_paid=200)).warnings)
    with pytest.raises(ValidationError):
        fee(payment_date=TODAY - dt.timedelta(days=20))  # before the PIO's letter


# ----------------------------------------------------------------------------- rendering
def test_docx_is_valid_and_uses_a_devanagari_font_for_marathi():
    doc = templates.build(rti(language="mr", full_name="रमेश कुलकर्णी"))
    word = Document(io.BytesIO(render.to_docx(doc)))
    body = "\n".join(p.text for p in word.paragraphs) + "\n".join(c.text for t in word.tables for r in t.rows for c in r.cells)
    assert "रमेश कुलकर्णी" in body and "अर्जदाराचे संपूर्ण नाव" in body
    run = next(r for p in word.paragraphs for r in p.runs if r.text)
    assert run._element.rPr.rFonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cs") == "Mangal"
    en = Document(io.BytesIO(render.to_docx(templates.build(rti()))))
    assert any("Ramesh Kulkarni" in c.text for t in en.tables for r in t.rows for c in r.cells)


def test_pdf_has_the_text_fits_the_page_and_escapes_html():
    doc = templates.build(rti(subject="Tom & Jerry <b>bold</b> plan"))
    pdf = pymupdf.open(stream=render.to_pdf(doc), filetype="pdf")
    page_text = "".join(p.get_text() for p in pdf)
    assert pdf.page_count <= 2 and "Ramesh Kulkarni" in page_text and "Tom & Jerry <b>bold</b> plan" in page_text
    right_edge = max(w[2] for w in pdf[0].get_text("words"))
    assert right_edge <= 595 - 40  # nothing is clipped at the right margin


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract not installed")
def test_marathi_pdf_renders_readable_devanagari():
    import pytesseract
    from PIL import Image

    pdf = pymupdf.open(stream=render.to_pdf(templates.build(rti(language="mr", full_name="रमेश कुलकर्णी"))), filetype="pdf")
    img = Image.open(io.BytesIO(pdf[0].get_pixmap(dpi=200).tobytes("png")))
    ocr = pytesseract.image_to_string(img, lang="mar+eng")
    assert "अर्जदाराचे" in ocr and "माहितीचा अधिकार" in ocr  # correctly shaped glyphs, not boxes or broken conjuncts


# ----------------------------------------------------------------------------- API
@pytest.fixture
def client(container):
    app.state.container = container
    with TestClient(app) as c:
        yield c


PAYLOAD = {"kind": "rti_application", "language": "en", "full_name": "Ramesh Kulkarni", "address": "Kothrud, Pune", "place": "Pune",
           "authority": "Pune Municipal Corporation", "subject": "Road repair", "items": ["Work orders for Ward 12"]}


def test_draft_endpoints_need_the_internal_token(client):
    assert client.post("/v1/drafts/preview", json=PAYLOAD).status_code == 401
    assert client.post("/v1/drafts/file", json=PAYLOAD).status_code == 401


def test_draft_preview_and_files(client):
    r = client.post("/v1/drafts/preview", json=PAYLOAD, headers=H)
    assert r.status_code == 200 and r.json()["title"] == "RTI application" and r.json()["blocks"]
    pdf = client.post("/v1/drafts/file?format=pdf", json=PAYLOAD, headers=H)
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF")
    assert 'filename="RTI-application-' in pdf.headers["content-disposition"]
    docx = client.post("/v1/drafts/file?format=docx", json=PAYLOAD, headers=H)
    assert docx.status_code == 200 and docx.content[:2] == b"PK" and docx.headers["content-type"] == render.DOCX_MIME
    assert client.post("/v1/drafts/file?format=exe", json=PAYLOAD, headers=H).status_code == 422


def test_draft_validation_errors_are_readable(client):
    bad = {**PAYLOAD, "fee_mode": "cash"}
    r = client.post("/v1/drafts/preview", json=bad, headers=H)
    assert r.status_code == 422 and r.json()["detail"] == "fee_reference: is required when the fee is paid by cash, demand draft or banker's cheque"
    assert r.json()["code"] == "validation"
    r = client.post("/v1/drafts/preview", json={**PAYLOAD, "items": []}, headers=H)
    assert r.status_code == 422 and "items" in r.json()["detail"]
    assert client.post("/v1/drafts/preview", json={"kind": "nonsense"}, headers=H).status_code == 422
    # other endpoints keep FastAPI's standard 422 shape
    assert isinstance(client.post("/v1/simplify", json={}, headers=H).json()["detail"], list)


def test_draft_request_union_picks_the_model_by_kind():
    ta = TypeAdapter(DraftRequest)
    assert isinstance(ta.validate_python({**PAYLOAD}), RtiApplication)


def appeal_payload(**kw):
    return {"kind": "first_appeal", "language": "en", "full_name": "R", "address": "Pune", "place": "Pune", "appellate_authority": "FAA",
            "pio_details": "PIO", "application_date": (TODAY - dt.timedelta(days=45)).isoformat(), "grounds": ["no_reply"],
            "info_subject": "x", "info_office": "y", **kw}


def fee_payload(**kw):
    return {"kind": "fee_payment_reply", "language": "mr", "full_name": "R", "address": "Pune", "place": "Pune", "pio_office": "PMC",
            "pio_letter_date": (TODAY - dt.timedelta(days=10)).isoformat(), "application_date": (TODAY - dt.timedelta(days=30)).isoformat(),
            "amount_demanded": 148, "pages": 74, "amount_paid": 148, "payment_mode": "demand_draft", "payment_reference": "DD1",
            "payment_date": (TODAY - dt.timedelta(days=5)).isoformat(), **kw}


def test_every_draft_kind_works_through_the_api_in_both_languages(client):
    for payload in (PAYLOAD, appeal_payload(), appeal_payload(language="mr", order_received_date=(TODAY - dt.timedelta(days=3)).isoformat(), grounds=["incomplete"]),
                    fee_payload(), fee_payload(language="en", amount_demanded=300, amount_paid=300)):
        r = client.post("/v1/drafts/preview", json=payload, headers=H)
        assert r.status_code == 200, (payload["kind"], r.text)
        for fmt in ("pdf", "docx"):
            f = client.post(f"/v1/drafts/file?format={fmt}", json=payload, headers=H)
            assert f.status_code == 200 and len(f.content) > 1000, (payload["kind"], fmt)


def test_validation_errors_mention_only_the_chosen_kind(client):
    """A form of one kind must not be reported against the other kinds (no 'RtiApplication.authority: Field required')."""
    r = client.post("/v1/drafts/preview", json=appeal_payload(application_date=(TODAY + dt.timedelta(days=40)).isoformat()), headers=H)
    assert r.status_code == 422
    assert r.json()["detail"] == "application_date: is in the future - check that the day and month are the right way round"
    r = client.post("/v1/drafts/preview", json={"kind": "fee_payment_reply", "language": "en", "full_name": "R", "address": "Pune", "place": "Pune"}, headers=H)
    detail = r.json()["detail"]
    assert "pio_office: is required" in detail and "amount_paid: is required" in detail
    assert "authority" not in detail.replace("pio_office", "") and "function-after" not in detail and "Field required" not in detail
    assert client.post("/v1/drafts/preview", json={**appeal_payload(), "application_date": "31/12/2026"}, headers=H).json()["detail"] == "application_date: is not a valid date"
    assert "kind:" in client.post("/v1/drafts/preview", json={"kind": "nope"}, headers=H).json()["detail"]


def test_dates_up_to_tomorrow_are_accepted_for_time_zone_differences(client):
    tomorrow = (TODAY + dt.timedelta(days=1)).isoformat()
    assert client.post("/v1/drafts/preview", json=appeal_payload(application_date=tomorrow), headers=H).status_code == 200
    assert client.post("/v1/drafts/preview", json=fee_payload(payment_date=tomorrow), headers=H).status_code == 200


def test_fee_reply_date_order_messages_are_plain(client):
    r = client.post("/v1/drafts/preview", json=fee_payload(payment_date=(TODAY - dt.timedelta(days=20)).isoformat()), headers=H)
    assert r.json()["detail"] == "payment_date: the payment cannot be dated before the PIO's letter"
    r = client.post("/v1/drafts/preview", json=fee_payload(pio_letter_date=(TODAY - dt.timedelta(days=40)).isoformat(), payment_date=(TODAY - dt.timedelta(days=39)).isoformat()), headers=H)
    assert r.json()["detail"] == "pio_letter_date: the PIO's letter cannot be dated before your application"
